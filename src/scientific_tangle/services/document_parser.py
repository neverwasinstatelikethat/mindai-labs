from __future__ import annotations

import io
import json
from pathlib import Path
from typing import Literal

import fitz
import pytesseract
from docx import Document
from openpyxl import load_workbook
from PIL import Image

from scientific_tangle.domain.contracts import DocumentFragment, DocumentRequest

MAX_DOCUMENT_BYTES = 20 * 1024 * 1024
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".json", ".txt"}


class UnsupportedDocumentError(ValueError):
    pass


def parse_document(
    filename: str,
    content: bytes,
    *,
    language: Literal["ru", "en"] = "ru",
    geography: str | None = None,
    year: int | None = None,
    ocr: bool = True,
) -> DocumentRequest:
    if not content:
        raise UnsupportedDocumentError("Файл пуст")
    if len(content) > MAX_DOCUMENT_BYTES:
        raise UnsupportedDocumentError("Файл превышает лимит 20 МБ")

    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise UnsupportedDocumentError(f"Формат {suffix or 'без расширения'} не поддерживается")

    fragments = _parse_pdf(content, ocr=ocr) if suffix == ".pdf" else _PARSERS[suffix](content)
    text = "\n\n".join(fragment.text for fragment in fragments if fragment.text.strip())
    if len(text.strip()) < 20:
        raise UnsupportedDocumentError("Не удалось извлечь достаточно текста")
    return DocumentRequest(
        title=Path(filename).stem,
        text=text,
        language=language,
        geography=geography,
        year=year,
        fragments=fragments,
    )


def _parse_pdf(content: bytes, *, ocr: bool = True) -> list[DocumentFragment]:
    rendered = fitz.open(stream=content, filetype="pdf")
    fragments: list[DocumentFragment] = []
    for index, page in enumerate(rendered, start=1):
        text = page.get_text("text").strip()
        if not text and ocr:
            pixmap = page.get_pixmap(dpi=200, alpha=False)
            image = Image.open(io.BytesIO(pixmap.tobytes("png")))
            text = pytesseract.image_to_string(image, lang="rus+eng").strip()
        if text:
            fragments.append(DocumentFragment(text=text, page=index))
    return fragments


def _parse_docx(content: bytes) -> list[DocumentFragment]:
    document = Document(io.BytesIO(content))
    paragraphs = [
        paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()
    ]
    fragments: list[DocumentFragment] = []
    block: list[str] = []
    block_size = 0
    block_start = 1
    for index, paragraph in enumerate(paragraphs, start=1):
        if block and block_size + len(paragraph) > 8000:
            fragments.append(
                DocumentFragment(
                    text="\n".join(block),
                    page=1,
                    sheet="body",
                    cell_range=f"paragraphs:{block_start}-{index - 1}",
                )
            )
            block = []
            block_size = 0
            block_start = index
        block.append(paragraph)
        block_size += len(paragraph)
    if block:
        fragments.append(
            DocumentFragment(
                text="\n".join(block),
                page=1,
                sheet="body",
                cell_range=f"paragraphs:{block_start}-{len(paragraphs)}",
            )
        )
    for table_index, table in enumerate(document.tables, start=1):
        for row_index, row in enumerate(table.rows, start=1):
            values = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if values:
                fragments.append(
                    DocumentFragment(
                        text=" | ".join(values),
                        sheet=f"table:{table_index}",
                        cell_range=f"row:{row_index}",
                    )
                )
    return fragments


def _parse_xlsx(content: bytes) -> list[DocumentFragment]:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    fragments: list[DocumentFragment] = []
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            filled_cells = [cell for cell in row if cell.value is not None]
            values = [str(cell.value) for cell in filled_cells]
            if filled_cells:
                fragments.append(
                    DocumentFragment(
                        text=" | ".join(values),
                        sheet=worksheet.title,
                        cell_range=(f"{filled_cells[0].coordinate}:{filled_cells[-1].coordinate}"),
                    )
                )
    return fragments


def _parse_json(content: bytes) -> list[DocumentFragment]:
    payload = json.loads(content.decode("utf-8-sig"))
    text = json.dumps(payload, ensure_ascii=False, indent=2)
    return [DocumentFragment(text=text, page=1)]


def _parse_text(content: bytes) -> list[DocumentFragment]:
    return [DocumentFragment(text=content.decode("utf-8-sig"), page=1)]


_PARSERS = {
    ".docx": _parse_docx,
    ".xlsx": _parse_xlsx,
    ".json": _parse_json,
    ".txt": _parse_text,
}
